import io

from openpyxl import load_workbook

from bidbrief.sample import ALTERNATES, BID_ITEMS
from bidbrief.schemas import BidItem
from bidbrief.tables import extract_bid_items, is_unit, items_to_csv, items_to_excel, numbering_gaps


def test_bid_items_match_the_form_exactly(ingested):
    res = extract_bid_items(ingested, use_model=False)
    got = [(i.item_no, i.description, i.quantity, i.unit) for i in res.items]
    want = [(a, d, q, u) for a, d, q, u in BID_ITEMS + ALTERNATES]
    assert got == want
    assert len(res.items) == 26
    assert res.numbering_gaps == [] and res.duplicates == []
    assert all(i.confidence >= 0.8 for i in res.items)
    assert {p for _, p in res.bid_form_pages} == {6, 7, 8}


def test_quantities_are_kept_as_printed_in_excel(ingested):
    res = extract_bid_items(ingested, use_model=False)
    wb = load_workbook(io.BytesIO(items_to_excel(res.items)))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=5, values_only=True))
    assert len(rows) == 26
    qty = {r[0]: r[2] for r in rows}
    assert qty["4"] == "3,120" and isinstance(qty["4"], str)
    assert qty["2"] == "60,000"
    csv_text = items_to_csv(res.items).decode()
    assert "3,120" in csv_text.replace('"', "")


def test_numbering_gaps_found():
    items = [BidItem(item_no=n, description="x", quantity="1", unit="EA", page=1, confidence=1) for n in
             ["1", "2", "4", "5", "100", "101"]]
    assert numbering_gaps(items) == ["3"]


def test_units():
    for u in ["LF", "L.F.", "EA", "LS", "Lump Sum", "CY", "TON", "$"]:
        assert is_unit(u), u
    assert not is_unit("Dollars and cents")
